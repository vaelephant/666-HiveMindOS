import hashlib
import uuid
from pathlib import Path

from memory_layer.knowledge_base import config
from memory_layer.knowledge_base.app.logging_config import get_logger

log = get_logger("hivemind.agent.ingest")
from memory_layer.knowledge_base.core.compiler.entity_extractor import extract_entities
from memory_layer.knowledge_base.core.compiler.workflow_extractor import extract_workflows
from memory_layer.knowledge_base.core.compiler.markdown_writer import (
    write_entity_page, write_workflow_page, write_rule_page,
)
from memory_layer.knowledge_base.core.wiki.wiki_manager import WikiManager
from memory_layer.knowledge_base.core.graph.memory_graph import MemoryGraph
from memory_layer.knowledge_base.models.entity import Entity

_TYPE_MAP = {".pdf": "pdf", ".docx": "word", ".doc": "word", ".xlsx": "excel", ".xls": "excel"}


class IngestAgent:
    def __init__(self, wiki: WikiManager, graph: MemoryGraph):
        self.wiki = wiki
        self.graph = graph

    def run(self, file_path: Path, org_id: str, source_type: str = "pdf") -> dict:
        log.info("extracting text  file=%s  type=%s", file_path.name, source_type)
        content = self._extract_text(file_path, source_type)
        content_hash = hashlib.sha256(content.encode()).hexdigest()
        log.info("text extracted   chars=%d  hash=%s…", len(content), content_hash[:8])

        log.info("extracting entities via LLM…")
        entities = extract_entities(content)
        log.info("entities found   count=%d", len(entities))

        log.info("extracting workflows via LLM…")
        workflows_data = extract_workflows(content)
        log.info("workflows found  count=%d  rules=%d",
                 len(workflows_data.get("workflows", [])),
                 len(workflows_data.get("rules", [])))

        pages: list[str] = []

        for e in entities:
            wiki_path = write_entity_page(self.wiki.root, org_id, e)
            pages.append(wiki_path)
            self.graph.upsert_entity(Entity(
                id=str(uuid.uuid4()),
                org_id=org_id,
                name=e["name"],
                entity_type=e["type"],
                wiki_path=wiki_path,
                attributes=e.get("attributes", {}),
            ))

        for wf in workflows_data.get("workflows", []):
            pages.append(write_workflow_page(self.wiki.root, org_id, wf))

        for rule in workflows_data.get("rules", []):
            pages.append(write_rule_page(self.wiki.root, org_id, rule))

        self.wiki.update_index(org_id)
        log.info("wiki index updated  org=%s  total_pages=%d", org_id, len(pages))

        return {
            "content_hash": content_hash,
            "entities_extracted": len(entities),
            "workflows_extracted": len(workflows_data.get("workflows", [])),
            "rules_extracted": len(workflows_data.get("rules", [])),
            "wiki_pages_created": len(pages),
            "pages": pages,
        }

    def _extract_text(self, file_path: Path, source_type: str) -> str:
        extractors = {
            "pdf": self._pdf,
            "word": self._word,
            "excel": self._excel,
        }
        return extractors.get(source_type, self._plain)(file_path)

    def _pdf(self, path: Path) -> str:
        try:
            import PyPDF2
            reader = PyPDF2.PdfReader(str(path))
            return "\n".join(p.extract_text() or "" for p in reader.pages)
        except Exception:
            return ""

    def _word(self, path: Path) -> str:
        try:
            import docx
            return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
        except Exception:
            return ""

    def _excel(self, path: Path) -> str:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(str(path), data_only=True)
            rows = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    rows.append("\t".join(str(c) for c in row if c is not None))
            return "\n".join(rows)
        except Exception:
            return ""

    def _plain(self, path: Path) -> str:
        return path.read_text(encoding="utf-8", errors="ignore")
